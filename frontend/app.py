import streamlit as st
import requests
import re
from io import BytesIO

st.set_page_config(page_title="LegalEase", layout="wide")
st.title("⚖️ LegalEase - AI Legal Document Generator")

def sanitize(text):
    if not text:
        return ""
    text = re.sub(r'[^\x00-\x7F]+', ' ', text)
    text = text.replace('\r','').replace('\x00','')
    text = re.sub(r'_{4,}', '________', text)
    text = re.sub(r'\*{4,}', '****', text)
    return text

def format_pdf(text, doc_type):
    from fpdf import FPDF
    pdf = FPDF(format='A4')
    pdf.set_margins(15, 15, 15)
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, doc_type.upper()[:80], align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(5)
    pdf.set_font("Helvetica", "", 10)
    clean = sanitize(text) or "No content"
    pdf.multi_cell(w=0, h=6, text=clean)
    return pdf

def create_docx(text):
    from docx import Document
    doc = Document()
    doc.add_heading("Legal Document", 1)
    for para in text.split("\n"):
        if para.strip():
            doc.add_paragraph(para)
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()

# ---------- UI INPUTS ----------
document_type = st.selectbox("Document Type", ["Employment Agreement", "NDA", "Service Agreement", "Rental Agreement"])
parties = st.text_area("Parties", "Amit Sharma (Employee) and XYZ Company (Employer)")
terms = st.text_area("Terms (separate by ;)", "Salary 50000 per month; Work from office; 30 days notice period")
effective_date = st.text_input("Effective Date", "27-09-2026")

if st.button("Generate Document", type="primary"):
    try:
        with st.spinner("Generating... Make sure backend is running on 8000"):
            res = requests.post("http://127.0.0.1:8000/generate", json={
                "document_type": document_type,
                "parties": parties,
                "terms": terms,
                "effective_date": effective_date,
                "dates": effective_date
            }, timeout=90)
            data = res.json()
            if "document" in data:
                st.session_state["doc"] = data["document"]
                st.success("Generated!")
            else:
                st.error(f"Backend error: {data}")
    except Exception as e:
        st.error(f"Error: {e}. Is backend running? Run: uvicorn backend.main:app --reload --port 8000")

# ---------- EDIT AND DOWNLOAD SECTION ----------
if "doc" in st.session_state:
    st.divider()
    st.subheader("📝 Edit Document")
    edited = st.text_area("You can edit before download", st.session_state["doc"], height=400, key="edit_area")
    st.session_state["doc"] = edited

    col1, col2, col3 = st.columns(3)

    # TXT
    with col1:
        st.download_button(
            "⬇️ Download TXT",
            data=edited.encode('utf-8'),
            file_name=f"{document_type}.txt",
            mime="text/plain",
            use_container_width=True
        )

    # DOCX
    with col2:
        try:
            docx_bytes = create_docx(edited)
            st.download_button(
                "⬇️ Download DOCX",
                data=docx_bytes,
                file_name=f"{document_type}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True
            )
        except Exception as e:
            st.error(f"Install python-docx: pip install python-docx {e}")

    # PDF
    with col3:
        try:
            pdf = format_pdf(edited, document_type)
            out = pdf.output(dest='S')
            pdf_bytes = out if isinstance(out, (bytes, bytearray)) else out.encode('latin1', errors='ignore')
            st.download_button(
                "⬇️ Download PDF",
                data=bytes(pdf_bytes),
                file_name=f"{document_type}.pdf",
                mime="application/pdf",
                use_container_width=True
            )
        except Exception as e:
            st.error(f"PDF Error: {e}")